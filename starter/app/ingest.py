"""Crash-proof corpus walk and per-type parsers for MC-3.

The corpus deliberately contains an unreadable file, an encrypted PDF, an
unknown binary and an empty directory. Nothing in here may let one bad file
stop the walk: every file is parsed in its own child process, under its own
try/except and its own timeout, and its outcome is recorded in the manifest.

Text formats are parsed here. Images (and PDF pages with no text layer) are
returned as OCR jobs and transcribed afterwards by vlm.py on the GPU -- that
pass runs in the parent process only after every child has exited, so we never
fork after the GPU runtime is initialised.
"""

from __future__ import annotations

import csv
import io
import json
import multiprocessing as mp
import os
import re
import stat
import sys
import time
import zipfile
from collections import Counter, deque
from pathlib import Path
from typing import Callable, Iterator

MAX_FILE_BYTES = 200 * 1024 * 1024
FILE_TIMEOUT_S = 30.0
CHUNK_CHARS = 800
MIN_PAGE_TEXT = 20  # a PDF page with less text than this is treated as scanned

TEXT_EXTS = {
    ".txt", ".log", ".md", ".rst", ".py", ".js", ".ts", ".java", ".c", ".h",
    ".cpp", ".go", ".rs", ".sh", ".json", ".yaml", ".yml", ".toml", ".ini",
    ".cfg", ".conf", ".xml", ".html", ".htm", ".sql",
}
CODE_EXTS = {".py", ".js", ".ts", ".java", ".c", ".h", ".cpp", ".go", ".rs", ".sh"}
PROSE_EXTS = {".txt", ".md", ".rst"}

SUPERSEDED_RE = re.compile(r"\b(WITHDRAWN|SUPERSEDED\s+BY|OBSOLETE|DEPRECATED)\b", re.I)
REV_NAME_RE = re.compile(r"(?:^|[^a-z0-9])(?:rev|r)[-_ ]?([0-9]+|[a-z][0-9]*)(?=[^a-z0-9]|$)")
REV_TEXT_RE = re.compile(r"\brevision\s+([A-Za-z0-9.-]+)", re.I)
ASSIGN_RE = re.compile(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s*(?::[^=]+)?=\s*([^=].*)$")


# --------------------------------------------------------------------------
# walk
# --------------------------------------------------------------------------

def walk(root: Path, dir_errors: list[dict]) -> Iterator[tuple[str, Path]]:
    """Yield (corpus-relative posix path, absolute path) for every regular file.

    Deterministic order, symlinks skipped, unreadable directories recorded.
    """
    def onerror(err: OSError) -> None:
        dir_errors.append({"path": _rel(root, Path(err.filename or root)),
                           "detail": f"{type(err).__name__}: {err.strerror or err}"})

    for dirpath, dirnames, filenames in os.walk(root, onerror=onerror, followlinks=False):
        dirnames[:] = sorted(d for d in dirnames if not os.path.islink(os.path.join(dirpath, d)))
        for name in sorted(filenames):
            abspath = Path(dirpath) / name
            try:
                st = os.lstat(abspath)
            except OSError as e:
                dir_errors.append({"path": _rel(root, abspath), "detail": f"lstat: {e}"})
                continue
            if stat.S_ISREG(st.st_mode):
                yield _rel(root, abspath), abspath


def _rel(root: Path, p: Path) -> str:
    try:
        return p.relative_to(root).as_posix()
    except ValueError:
        return p.as_posix()


# --------------------------------------------------------------------------
# sniffing
# --------------------------------------------------------------------------

def sniff(abspath: Path, head: bytes) -> str:
    """Type from magic bytes first, then extension. Never libmagic: it calls
    arbitrary binary blobs 'Targa image data'."""
    ext = abspath.suffix.lower()
    if b"%PDF" in head[:1024]:
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        with zipfile.ZipFile(abspath) as z:
            names = set(z.namelist())
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        return "unknown_binary"
    if head.startswith((b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"GIF87a", b"GIF89a", b"BM")) \
            or (head[:4] == b"RIFF" and head[8:12] == b"WEBP") \
            or head[:4] in (b"II*\x00", b"MM\x00*"):
        return "image"
    if not _looks_like_text(head):
        return "unknown_binary"
    if ext in (".csv", ".tsv"):
        return "csv"
    return "text"


def _looks_like_text(head: bytes) -> bool:
    if b"\x00" in head:
        return False
    try:
        # final=False: a multi-byte character cut off at the sample boundary is fine
        import codecs
        codecs.getincrementaldecoder("utf-8")().decode(head, final=False)
        return True
    except UnicodeDecodeError:
        printable = sum(32 <= b < 127 or b in (9, 10, 13) for b in head)
        return bool(head) and printable / len(head) > 0.95


# --------------------------------------------------------------------------
# per-file processing (runs in a child process)
# --------------------------------------------------------------------------

def process_file(abspath: Path, rel: str, max_bytes: int = MAX_FILE_BYTES) -> dict:
    """Parse one file. Returns {"record", "chunks", "ocr_jobs"}; never raises
    for an expected failure, and the caller guards the unexpected ones."""
    rec = {"path": rel, "type": None, "status": "ok", "n_chunks": 0, "detail": ""}
    out = {"record": rec, "chunks": [], "ocr_jobs": []}

    # Readability is decided by actually opening the file, not os.access():
    # with DAC_OVERRIDE dropped, root's os.access answer is not what open() does.
    try:
        size = os.stat(abspath).st_size
        with open(abspath, "rb") as f:
            head = f.read(4096)
    except OSError as e:
        rec.update(status="unreadable", detail=f"{type(e).__name__}: {e.strerror or e}")
        return out

    if size == 0:
        rec.update(status="empty")
        return out
    if size > max_bytes:
        rec.update(status="too_large", detail=f"{size} bytes")
        return out

    kind = sniff(abspath, head)
    rec["type"] = kind
    if kind == "unknown_binary":
        rec.update(status="unknown_binary", detail=f"head={head[:8].hex()}")
        return out

    parser = PARSERS[kind]
    status, detail, chunks, jobs = parser(abspath, rel)
    meta = _file_meta(rel, kind, chunks)
    for i, c in enumerate(chunks):
        c["meta"] = {**meta, **c.get("meta", {})}
        c["id"] = f"{rel}#{i}"
    for j in jobs:
        j["meta"] = meta

    rec.update(status=status, detail=detail, n_chunks=len(chunks))
    if status == "ok" and not chunks and not jobs:
        rec["status"] = "empty"
    out["chunks"], out["ocr_jobs"] = chunks, jobs
    return out


def _chunk(rel: str, kind: str, locator: str, text: str, **meta) -> dict:
    return {"path": rel, "type": kind, "locator": locator, "text": text, "meta": meta}


def _file_meta(rel: str, kind: str, chunks: list[dict]) -> dict:
    """Supersession and revision flags, used by retrieval to prefer the
    current document. Content is only checked for document-like files: a CSV
    row or a log line saying 'deprecated' must not demote the whole file."""
    name = Path(rel).name
    meta: dict = {}
    head = ""
    if kind in ("pdf", "docx") or (kind == "text" and Path(rel).suffix.lower() in PROSE_EXTS):
        head = " ".join(c["text"] for c in chunks[:2])[:400]
    meta["superseded"] = bool(SUPERSEDED_RE.search(name) or SUPERSEDED_RE.search(head))
    m = REV_NAME_RE.search(Path(rel).stem.lower()) or REV_TEXT_RE.search(head)
    if m:
        meta["revision"] = m.group(1)
    return meta


# ---- PDF -----------------------------------------------------------------

def parse_pdf(abspath: Path, rel: str):
    # Raw check before pypdf: an AES-256 PDF can need a crypto backend just to
    # be constructed, and an empty user password would let pypdf decrypt it
    # silently. Either way it is content we must not answer from.
    if _pdf_has_encrypt_dict(abspath):
        return "encrypted", "/Encrypt in trailer", [], []

    from pypdf import PdfReader
    reader = PdfReader(str(abspath), strict=False)
    if reader.is_encrypted:
        return "encrypted", "is_encrypted", [], []

    chunks, jobs, errors = [], [], []
    for i, page in enumerate(reader.pages):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as e:  # one bad page does not lose the document
            errors.append(f"page {i + 1}: {type(e).__name__}")
            text = ""
        if len(text) < MIN_PAGE_TEXT:
            jobs.append({"path": rel, "kind": "pdf_page", "page": i, "locator": f"page {i + 1}"})
            continue
        for part in _split(text):
            chunks.append(_chunk(rel, "pdf", f"page {i + 1}", part))
    return "ok", "; ".join(errors), chunks, jobs


def _pdf_has_encrypt_dict(abspath: Path) -> bool:
    size = os.path.getsize(abspath)
    with open(abspath, "rb") as f:
        head = f.read(65536)
        f.seek(max(0, size - 65536))
        tail = f.read()
    return re.search(rb"/Encrypt\s*(\d+\s+\d+\s+R|<<)", head + tail) is not None


# ---- DOCX ----------------------------------------------------------------

def parse_docx(abspath: Path, rel: str):
    import docx
    doc = docx.Document(str(abspath))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for section in doc.sections:
        for part in (section.header, section.footer):
            try:
                paras += [p.text.strip() for p in part.paragraphs if p.text.strip()]
            except Exception:
                pass

    chunks = [_chunk(rel, "docx", "body", t) for t in _pack(paras)]
    for ti, table in enumerate(doc.tables):
        rows = [[c.text.strip() for c in r.cells] for r in table.rows]
        for ri, line in _table_rows(rows):
            chunks.append(_chunk(rel, "docx", f"table {ti + 1} row {ri + 1}", line))
    return "ok", "", chunks, []


# ---- XLSX ----------------------------------------------------------------

def parse_xlsx(abspath: Path, rel: str):
    import openpyxl
    wb = openpyxl.load_workbook(str(abspath), read_only=True, data_only=True)
    chunks = []
    try:
        for ws in wb.worksheets:  # every sheet: answers hide on the second one too
            rows = [["" if v is None else _cell(v) for v in row]
                    for row in ws.iter_rows(values_only=True)]
            for ri, line in _table_rows(rows):
                chunks.append(_chunk(rel, "xlsx", f"sheet {ws.title} row {ri + 1}",
                                     f"sheet={ws.title}; {line}", sheet=ws.title))
    finally:
        wb.close()
    return "ok", "", chunks, []


def _cell(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---- CSV -----------------------------------------------------------------

def parse_csv(abspath: Path, rel: str):
    raw = abspath.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if abspath.suffix.lower() == ".tsv" else csv.excel
    csv.field_size_limit(min(sys.maxsize, 2**31 - 1))
    rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
    chunks = [_chunk(rel, "csv", f"row {ri + 1}", line) for ri, line in _table_rows(rows)]
    return "ok", "", chunks, []


def _table_rows(rows: list[list[str]]) -> Iterator[tuple[int, str]]:
    """Yield (row index, 'header=value; ...') so a key and its values stay in
    one chunk. Rows before the header, or with a single cell, are free text."""
    header = None
    for ri, row in enumerate(rows):
        cells = [str(c).strip() for c in row]
        filled = [c for c in cells if c]
        if not filled:
            continue
        if header is None and len(filled) >= 2:
            header = [c or f"col{i + 1}" for i, c in enumerate(cells)]
            continue
        if header is None or len(filled) == 1:
            yield ri, " | ".join(filled)
            continue
        pairs = []
        for i, c in enumerate(cells):
            if c:
                key = header[i] if i < len(header) else f"col{i + 1}"
                pairs.append(f"{key}={c}")
        yield ri, "; ".join(pairs)


# ---- text / log / code ---------------------------------------------------

def parse_text(abspath: Path, rel: str):
    text = abspath.read_bytes().decode("utf-8", errors="replace")
    is_code = abspath.suffix.lower() in CODE_EXTS
    chunks = []
    for start, part in _split_lines(text):
        meta = {}
        if is_code:
            # Constants are tagged so the answer step can prefer them over a
            # comment that names an old value ("raised from 60").
            assigns = [f"{m.group(1)}={m.group(2).strip()}"
                       for line in part.splitlines()
                       if not line.lstrip().startswith("#") and (m := ASSIGN_RE.match(line))]
            if assigns:
                meta["assignments"] = assigns
        chunks.append(_chunk(rel, "text", f"line {start}", part, **meta))
    return "ok", "", chunks, []


def _split_lines(text: str, size: int = CHUNK_CHARS, overlap_lines: int = 2):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        j, n = i, 0
        while j < len(lines) and (n == 0 or n + len(lines[j]) <= size):
            n += len(lines[j]) + 1
            j += 1
        part = "\n".join(lines[i:j]).strip()
        if part:
            yield i + 1, part
        if j >= len(lines):
            break
        i = max(i + 1, j - overlap_lines)


def _split(text: str, size: int = CHUNK_CHARS * 2) -> list[str]:
    return [p for _, p in _split_lines(text, size)] or [text]


def _pack(paras: list[str], size: int = CHUNK_CHARS) -> list[str]:
    out, buf = [], ""
    for p in paras:
        if buf and len(buf) + len(p) > size:
            out.append(buf)
            buf = ""
        buf = f"{buf}\n{p}" if buf else p
    if buf:
        out.append(buf)
    return out


# ---- images --------------------------------------------------------------

def parse_image(abspath: Path, rel: str):
    return "ok", "", [], [{"path": rel, "kind": "image", "locator": "image"}]


PARSERS: dict[str, Callable] = {
    "pdf": parse_pdf, "docx": parse_docx, "xlsx": parse_xlsx,
    "csv": parse_csv, "text": parse_text, "image": parse_image,
}


# --------------------------------------------------------------------------
# isolated execution: one child process per file, killed on timeout
# --------------------------------------------------------------------------

def _child(conn, abspath: str, rel: str, max_bytes: int) -> None:
    try:
        res = process_file(Path(abspath), rel, max_bytes)
    except Exception as e:
        res = _failed(rel, "error", f"{type(e).__name__}: {e}")
    try:
        conn.send(res)
    except Exception as e:  # e.g. unpicklable content
        conn.send(_failed(rel, "error", f"send: {type(e).__name__}: {e}"))
    finally:
        conn.close()


def _failed(rel: str, status: str, detail: str) -> dict:
    return {"record": {"path": rel, "type": None, "status": status, "n_chunks": 0,
                       "detail": detail[:500]},
            "chunks": [], "ocr_jobs": []}


def run_isolated(files: list[tuple[str, Path]], workers: int, timeout: float,
                 max_bytes: int = MAX_FILE_BYTES) -> list[dict]:
    """Parse every file in its own process. A crash, hang or segfault in one
    parser costs that file only. Results come back in walk order."""
    ctx = mp.get_context("fork" if sys.platform.startswith("linux") else "spawn")
    pending = deque(enumerate(files))
    running: dict = {}
    results: list = [None] * len(files)

    while pending or running:
        while pending and len(running) < workers:
            idx, (rel, abspath) = pending.popleft()
            parent, child = ctx.Pipe(duplex=False)
            proc = ctx.Process(target=_child, args=(child, str(abspath), rel, max_bytes), daemon=True)
            t0 = time.monotonic()
            proc.start()
            child.close()
            running[proc] = (idx, rel, parent, t0)

        progressed = False
        for proc, (idx, rel, conn, t0) in list(running.items()):
            res = None
            try:
                if conn.poll():  # read before join: a full pipe would deadlock the child
                    res = conn.recv()
                elif not proc.is_alive():
                    res = _failed(rel, "error", f"worker died, exit code {proc.exitcode}")
                elif time.monotonic() - t0 > timeout:
                    proc.kill()
                    res = _failed(rel, "timeout", f"exceeded {timeout:.0f}s")
            except Exception as e:  # EOFError if the child died mid-send
                res = _failed(rel, "error", f"{type(e).__name__}: {e}")
            if res is None:
                continue
            res["record"]["seconds"] = round(time.monotonic() - t0, 3)
            results[idx] = res
            proc.join(5)
            if proc.is_alive():
                proc.kill()
            conn.close()
            del running[proc]
            progressed = True
        if not progressed:
            time.sleep(0.01)
    return results


# --------------------------------------------------------------------------
# build
# --------------------------------------------------------------------------

def build(corpus: str | Path, out_dir: str | Path,
          transcriber_factory: Callable[[], Callable] | None = None,
          workers: int | None = None, timeout: float = FILE_TIMEOUT_S,
          ocr_deadline: float | None = None) -> dict:
    """Walk, parse, OCR, persist. Returns the manifest.

    `transcriber_factory` builds the GPU VLM; it is called only after every
    parser process has exited. If it is None, OCR jobs are recorded as
    'ocr_skipped' (useful on a dev box with no GPU). The text index is saved
    before the model is loaded, so a VLM failure never loses it. OCR stops
    starting new jobs once `ocr_deadline` seconds have passed since the start.
    """
    root = Path(corpus).resolve()
    out = Path(out_dir)
    if not root.is_dir():
        raise NotADirectoryError(f"corpus root {root} is not a directory")
    out.mkdir(parents=True, exist_ok=True)

    t_start = time.monotonic()
    dir_errors: list[dict] = []
    files = list(walk(root, dir_errors))
    workers = workers or min(8, os.cpu_count() or 2)
    results = run_isolated(files, workers, timeout)

    records = [r["record"] for r in results]
    chunks = [c for r in results for c in r["chunks"]]
    jobs = [j for r in results for j in r["ocr_jobs"]]
    by_path = {rec["path"]: rec for rec in records}

    for j in jobs:
        rec = by_path[j["path"]]
        rec["status"] = "ocr_skipped" if rec["type"] == "image" else rec["status"]
    _save(out, root, records, chunks, dir_errors, t_start)

    transcribe = None
    if jobs and transcriber_factory is not None:
        try:
            transcribe = transcriber_factory()
        except Exception as e:
            # Keep the text index: the text questions can still score.
            print(f"[index] OCR unavailable, images not indexed: {type(e).__name__}: {e}",
                  file=sys.stderr)
    if transcribe is not None:
        for j in jobs:
            rec = by_path[j["path"]]
            if ocr_deadline is not None and time.monotonic() - t_start > ocr_deadline:
                rec["detail"] = "OCR time budget exhausted"
                continue
            try:
                text = transcribe(root / j["path"], j).strip()
            except Exception as e:  # one bad image never stops the pass
                _mark_ocr_failed(rec, j, f"{type(e).__name__}: {e}")
                continue
            if not text:
                _mark_ocr_failed(rec, j, "empty transcription")
                continue
            new = [_chunk(j["path"], rec["type"], j["locator"], part,
                          **j.get("meta", {}), source="vlm") for part in _split(text)]
            base = rec["n_chunks"]
            for i, c in enumerate(new):
                c["id"] = f"{j['path']}#{base + i}"
            chunks.extend(new)
            rec["n_chunks"] += len(new)
            if rec["status"] in ("ocr_skipped", "empty"):
                rec["status"] = "ok"
        _save(out, root, records, chunks, dir_errors, t_start)

    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    print(f"[index] {manifest['summary_line']}", file=sys.stderr)
    for rec in records:
        if rec["status"] != "ok":
            print(f"[index]   {rec['status']:<14} {rec['path']}  {rec['detail']}", file=sys.stderr)
    return manifest


def _mark_ocr_failed(rec: dict, job: dict, detail: str) -> None:
    msg = f"{job['locator']}: {detail}"[:300]
    rec["detail"] = f"{rec['detail']}; {msg}" if rec["detail"] else msg
    if rec["type"] == "image":
        rec["status"] = "ocr_failed"


def _save(out: Path, root: Path, records, chunks, dir_errors, t_start) -> None:
    counts = Counter(r["status"] for r in records)
    line = f"{len(records)} files: " + ", ".join(f"{n} {s}" for s, n in counts.most_common()) \
        + f"; {len(chunks)} chunks; {len(dir_errors)} dir errors"
    manifest = {
        "corpus": str(root),
        "summary": dict(counts),
        "summary_line": line,
        "n_chunks": len(chunks),
        "seconds": round(time.monotonic() - t_start, 2),
        "files": records,
        "dir_errors": dir_errors,
    }
    _atomic_write(out / "chunks.jsonl",
                  "".join(json.dumps(c, ensure_ascii=False) + "\n" for c in chunks))
    _atomic_write(out / "manifest.json", json.dumps(manifest, ensure_ascii=False, indent=1))


def _atomic_write(path: Path, data: str) -> None:
    tmp = path.with_name(path.name + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(data)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)


if __name__ == "__main__":
    # Dev entry point: python3 ingest.py <corpus> <out_dir>   (no OCR)
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "index")
